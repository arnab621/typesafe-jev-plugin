"""
TypeSafe Jev dataset runner.

Reads input data, applies a solution signature to each row via the TypeSafe
API, and writes an Excel output file with input columns plus extracted answer
columns defined by the output_template.

Supported input formats:
    Tabular  — CSV (.csv) or Excel (.xlsx/.xlsm)
    Text     — single .txt or .md file (one row = the whole file)
    Text dir — directory of .txt/.md files (one row per file)

State template placeholders:
    ${col:Column Name}  — value from a named tabular column
    ${file}             — full text content of the current text file
    ${filename}         — name of the current text file (e.g. "resume.txt")

Usage:
    python run_dataset.py \\
        --input     path/to/data.xlsx \\
        --signature path/to/my.sig.json \\
        [--output   path/to/results.xlsx] \\
        [--api-key  YOUR_KEY]

API key precedence: --api-key flag > TYPESAFE_API_KEY environment variable.

Dependencies: requests, openpyxl
    pip install requests openpyxl
"""

import argparse
import csv
import json
import os
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("ERROR: openpyxl is required.  pip install openpyxl")
    sys.exit(1)

try:
    import requests
except ImportError:
    print("ERROR: requests is required.  pip install requests")
    sys.exit(1)

# Import the API client from the same directory
sys.path.insert(0, str(Path(__file__).parent))
from typesafe_client import call_api


# ---------------------------------------------------------------------------
# State building — expand ${col:Name}, ${file}, and ${filename} placeholders
# ---------------------------------------------------------------------------

def build_state(template, row):
    """Recursively expand state template placeholders using row values.

    Supported placeholders:
        ${col:Name}   — value of a named column (tabular mode)
        ${file}       — full text content of the input file (text mode)
        ${filename}   — name of the input file, e.g. "resume.txt" (text mode)
    """
    if isinstance(template, dict):
        return {k: build_state(v, row) for k, v in template.items()}
    elif isinstance(template, list):
        return [build_state(item, row) for item in template]
    elif isinstance(template, str):
        def replace_col(match):
            col_name = match.group(1)
            if col_name not in row:
                visible = ", ".join(f"'{c}'" for c in row.keys() if not c.startswith("__"))
                raise ValueError(f"Column '{col_name}' not found in input. Available: {visible}")
            val = row[col_name]
            return "" if val is None else str(val)

        result = re.sub(r'\$\{col:([^}]+)\}', replace_col, template)
        result = result.replace("${file}",     str(row.get("__content__",  "")))
        result = result.replace("${filename}", str(row.get("__filename__", "")))
        return result
    else:
        return template


# ---------------------------------------------------------------------------
# Output extraction — resolve dot-paths into the API response
# ---------------------------------------------------------------------------

def extract_path(data, path):
    """Resolve a dot-path (e.g. 'answers.vat_category.choice') in a nested dict."""
    current = data
    for part in path.split("."):
        if isinstance(current, dict) and part in current:
            current = current[part]
        else:
            return None
    return current


def extract_output_row(response, output_template):
    """Build output column dict from an API response using the output_template spec."""
    if not output_template or "columns" not in output_template:
        return {}

    result = {}
    for col_def in output_template["columns"]:
        header = col_def["header"]
        col_type = col_def.get("type", "path")

        if col_type == "confidence_flag":
            confidence = extract_path(response, col_def["path"])
            threshold = col_def.get("threshold", 0.6)
            if confidence is None:
                result[header] = "Unknown"
            else:
                result[header] = "Yes" if confidence < threshold else "No"

        else:  # default: dot-path extraction
            value = extract_path(response, col_def["path"])
            if isinstance(value, dict):
                value = json.dumps(value)   # serialise probability dicts as JSON string
            elif value is None:
                value = ""
            result[header] = value

    return result


# ---------------------------------------------------------------------------
# Input reading
# ---------------------------------------------------------------------------

def read_csv(path):
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(dict(row))
    return rows


def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in excel_row):
            continue  # skip empty rows
        rows.append({h: v for h, v in zip(headers, excel_row) if h is not None})
    return rows


def read_text_file(path):
    """Read a single .txt or .md file as one row (text analysis mode)."""
    p = Path(path)
    content = p.read_text(encoding="utf-8")
    return [{"__content__": content, "__filename__": p.name}]


def read_text_dir(path):
    """Read all .txt and .md files in a directory — one row per file."""
    p = Path(path)
    files = sorted(p.glob("*.txt")) + sorted(p.glob("*.md"))
    if not files:
        raise ValueError(f"No .txt or .md files found in directory: {path}")
    return [
        {"__content__": f.read_text(encoding="utf-8"), "__filename__": f.name}
        for f in files
    ]


def read_input(path):
    p = Path(path)
    if p.is_dir():
        return read_text_dir(path)
    suffix = p.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    elif suffix in (".xlsx", ".xls", ".xlsm"):
        return read_excel(path)
    elif suffix in (".txt", ".md"):
        return read_text_file(path)
    else:
        raise ValueError(
            f"Unsupported input '{suffix}'. Use .csv, .xlsx, .txt, .md, "
            "or a directory of .txt/.md files"
        )


# ---------------------------------------------------------------------------
# Output writing
# ---------------------------------------------------------------------------

HEADER_FILL   = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT   = Font(bold=True, color="FFFFFF")
FLAG_FILL     = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ERROR_FILL    = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")


def _display_row(row):
    """Convert an internal row to a display row: rename __filename__ → Source File, drop __content__."""
    out = {}
    if "__filename__" in row:
        out["Source File"] = row["__filename__"]
    for k, v in row.items():
        if not k.startswith("__"):
            out[k] = v
    return out


def write_results(input_rows, output_rows, output_path, sig_name, errors_by_index):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    if not input_rows:
        wb.save(output_path)
        return

    # Convert internal rows to display-friendly rows (strip __ private keys)
    display_input_rows = [_display_row(r) for r in input_rows]

    # Determine column layout
    input_headers  = list(display_input_rows[0].keys())
    output_headers = list(output_rows[0].keys()) if output_rows else []
    extra_headers  = [h for h in output_headers if h not in input_headers]
    all_headers    = input_headers + extra_headers

    # Header row
    ws.append(all_headers)
    for cell in ws[1]:
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for i, in_row in enumerate(display_input_rows):
        out_row   = output_rows[i] if i < len(output_rows) else {}
        row_data  = [in_row.get(h, "") for h in input_headers]
        row_data += [out_row.get(h, "") for h in extra_headers]
        ws.append(row_data)

        # Highlight error rows
        if i in errors_by_index:
            for cell in ws[ws.max_row]:
                cell.fill = ERROR_FILL

        # Highlight flagged rows (any "Needs Review" = Yes)
        elif any(str(out_row.get(h, "")).strip().lower() == "yes" for h in extra_headers):
            for cell in ws[ws.max_row]:
                cell.fill = FLAG_FILL

    # Auto-width columns
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Signature", sig_name])
    ws2.append(["Total rows",    len(input_rows)])
    ws2.append(["Succeeded",     len(input_rows) - len(errors_by_index)])
    ws2.append(["Errors",        len(errors_by_index)])
    flagged = sum(
        1 for out_row in output_rows
        if any(str(out_row.get(h, "")).strip().lower() == "yes" for h in extra_headers)
    )
    ws2.append(["Flagged (low confidence)", flagged])

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Run a TypeSafe Jev solution signature against a dataset"
    )
    parser.add_argument("--input",     required=True, help="Input: CSV/Excel file, single .txt/.md file, or directory of .txt/.md files")
    parser.add_argument("--signature", required=True, help="Path to .sig.json solution signature")
    parser.add_argument("--output",    default=None,  help="Output Excel path (default: <input>_results.xlsx)")
    parser.add_argument("--api-key",   default=None,  help="TypeSafe API key (default: TYPESAFE_API_KEY env var)")
    args = parser.parse_args()

    # Resolve API key
    api_key = args.api_key or os.environ.get("TYPESAFE_API_KEY")
    if not api_key:
        print("ERROR: TypeSafe API key required.")
        print("  Set with: export TYPESAFE_API_KEY=your-key-here")
        print("  Or use:   --api-key YOUR_KEY")
        sys.exit(1)

    # Load signature
    try:
        with open(args.signature, "r", encoding="utf-8") as f:
            sig = json.load(f)
    except FileNotFoundError:
        print(f"ERROR: Signature file not found: {args.signature}")
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"ERROR: Invalid JSON in signature: {e}")
        sys.exit(1)

    model           = sig.get("model", "jev-1.13.0")
    state_template  = sig.get("state_template", {})
    questions       = sig.get("questions", {})
    output_template = sig.get("output_template")

    # Load input data
    try:
        input_rows = read_input(args.input)
    except Exception as e:
        print(f"ERROR reading input file: {e}")
        sys.exit(1)

    if not input_rows:
        print("WARNING: Input file is empty — nothing to process.")
        sys.exit(0)

    # Resolve output path
    output_path = args.output
    if not output_path:
        p = Path(args.input)
        output_path = str(p.parent / f"{p.stem}_results.xlsx")

    print(f"Signature  : {sig.get('name', args.signature)}")
    print(f"Model      : {model}")
    print(f"Questions  : {len(questions)}")
    print(f"Input      : {args.input}  ({len(input_rows)} rows)")
    print(f"Output     : {output_path}")
    print()

    output_rows     = []
    errors_by_index = {}
    low_conf_rows   = []

    for i, row in enumerate(input_rows):
        row_num = i + 1
        print(f"  Row {row_num}/{len(input_rows)} ...", end=" ", flush=True)
        try:
            state    = build_state(state_template, row)
            response = call_api(state, questions, model, api_key)
            out_row  = extract_output_row(response, output_template)
            output_rows.append(out_row)

            # Detect low-confidence flags
            if output_template:
                for col_def in output_template.get("columns", []):
                    if col_def.get("type") == "confidence_flag":
                        conf      = extract_path(response, col_def["path"])
                        threshold = col_def.get("threshold", 0.6)
                        if conf is not None and conf < threshold:
                            low_conf_rows.append(row_num)
                            break

            print("OK")

        except requests.HTTPError as e:
            status = e.response.status_code if e.response is not None else "?"
            print(f"API ERROR {status}")
            errors_by_index[i] = f"HTTP {status}: {e}"
            output_rows.append({})

        except Exception as e:
            print(f"ERROR: {e}")
            errors_by_index[i] = str(e)
            output_rows.append({})

    # Write output
    print()
    try:
        write_results(
            input_rows, output_rows, output_path,
            sig.get("name", ""), errors_by_index
        )
    except Exception as e:
        print(f"ERROR writing output file: {e}")
        sys.exit(1)

    # Summary
    success = len(input_rows) - len(errors_by_index)
    print("=" * 52)
    print(f"  Rows processed : {len(input_rows)}")
    print(f"  Succeeded      : {success}")
    if low_conf_rows:
        print(f"  Flagged review : {len(low_conf_rows)}  (rows: {low_conf_rows})")
    if errors_by_index:
        print(f"  Errors         : {len(errors_by_index)}  (rows: {[i+1 for i in errors_by_index]})")
        for idx, msg in errors_by_index.items():
            print(f"    Row {idx+1}: {msg}")
    print(f"  Output         : {output_path}")
    print("=" * 52)

    sys.exit(0 if not errors_by_index else 1)


if __name__ == "__main__":
    main()
